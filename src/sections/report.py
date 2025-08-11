from docx import Document
from images import get_images_from_dir
from images import divide_images
from utils import substitute_placeholders, next_filename, search_paragraph
# from images import resize_images
from tables import create_non_conformities_table, create_town_units_table, create_documents_table

document = Document(r"../data/RELATÓRIO MODELO.docx")

def generate_report():
    # Busca a Sessão pra inserir as imagens
    search_non_conformity = search_paragraph(document,"APÊNDICE 1 - NÃO CONFORMIDADES")
    last_position = document.paragraphs[search_non_conformity[-1]]
    # resize_images()

    create_non_conformities_table(document, "Tabela 6 - Lista de NCs do SAA {{Municipio}}")
    create_town_units_table(document, "Tabela 2 - Descrição dos SAA {{Municipio}}.")
    create_documents_table(document, "Tabela 1 - Principais documentações solicitadas.")

    list_of_images_path = get_images_from_dir()
    divide_images(document, last_position, list_of_images_path)
    substitute_placeholders(document)
    
    print(">>> Alterações Concluidas")

    new_archive = next_filename()
    document.save(new_archive) 