from docx import Document
from images import get_images_from_dir
from images import divide_images
from utils import substitute_placeholders, next_filename, search_paragraph
# from images import resize_images
from tables import create_non_conformities_table, create_town_units_table, create_documents_table, create_statistics_table, create_quality_index_table, create_general_information_table, create_abbreviations_table, create_last_report_table

document = Document(r"./data/RELATÓRIO MODELO.docx")

def generate_report():
    # Busca a Sessão pra inserir as imagens
    search_non_conformity = search_paragraph(document,"APÊNDICE 1 - NÃO CONFORMIDADES")
    last_position = document.paragraphs[search_non_conformity[-1]]
    # resize_images()

    create_abbreviations_table(document, "LISTA DE ABREVIATURAS E SIGLAS")
    create_general_information_table(document, "3.	INFORMAÇÕES GERAIS")
    
    create_documents_table(document, "Tabela 1 - Principais documentações solicitadas.")
    create_town_units_table(document, "Tabela 2 - Descrição dos SAA {{Municipio}}.")
    create_last_report_table(document, "Tabela 3 - Contexto histórico resumido das fiscalizações do município de {{Municipio}}.")
    create_statistics_table(document, "Tabela 4 - Informações do prestador de serviços e do município de {{Municipio}}.")
    create_quality_index_table(document, "Tabela 5 - Principais Indicadores Regulatórios do município {{Municipio}}.")
    create_non_conformities_table(document, "Tabela 6 - Lista de NCs do SAA {{Municipio}}")

    list_of_images_path = get_images_from_dir()
    divide_images(document, last_position, list_of_images_path)
    substitute_placeholders(document)
    
    print(">>> Alterações Concluidas")

    new_archive = next_filename()
    document.save(new_archive) 